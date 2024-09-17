import tkinter as tk
from tkinter import scrolledtext, messagebox
from openpyxl import load_workbook
from prettytable import PrettyTable
import datetime

# Carregar a planilha Excel
tabela = load_workbook("GESTAO_DE_EXAMES_PERIODICOS.xlsx", data_only=True)
aba_ativa = tabela.active

# Função para formatar datas no padrão dd-mm-aaaa
def formatar_data(valor):
    if isinstance(valor, datetime.datetime):
        return valor.strftime('%d-%m-%Y')
    return valor

# Atualizar a lista de linhas formatadas e colunas válidas (excluindo "Coluna_10" até "Coluna_26")
def atualizar_dados():
    global linhas_formatadas, colunas_validas
    linhas_formatadas = []
    for linha in aba_ativa.iter_rows(min_row=5, max_col=9, values_only=True):  # Limitando as colunas até a 9ª (exclui Coluna_10 até Coluna_26)
        linha_formatada = [formatar_data(celula) for celula in linha]
        if any(celula is not None for celula in linha_formatada):
            linhas_formatadas.append(linha_formatada)
    
    # Ordenar as linhas formatadas pela primeira coluna
    linhas_formatadas.sort(key=lambda x: str(x[0]).lower() if x[0] is not None else "")

    colunas_validas = []
    for i in range(len(cabecalhos_unicos)):
        if i < 9:  # Limitar o número de colunas até 9
            if any(linha[i] is not None for linha in linhas_formatadas):
                colunas_validas.append(i)

# Definir cabeçalhos da tabela com base na 4ª linha (limitando até 9 colunas)
cabecalhos = [celula.value if celula.value is not None else f"Coluna_{i}" for i, celula in enumerate(aba_ativa[4][:9], start=1)]
cabecalhos_unicos = []
for i, cabecalho in enumerate(cabecalhos):
    if cabecalho in cabecalhos_unicos:
        cabecalhos_unicos.append(f"{cabecalho}_{i+1}")
    else:
        cabecalhos_unicos.append(cabecalho)

# Função para exibir a tabela na interface tkinter
def exibir_tabela():
    atualizar_dados()
    tabela_sem_colunas_vazias = PrettyTable()
    tabela_sem_colunas_vazias.field_names = [cabecalhos_unicos[i] for i in colunas_validas]
    for linha in linhas_formatadas:
        linha_filtrada = [linha[i] for i in colunas_validas]
        tabela_sem_colunas_vazias.add_row(linha_filtrada)

    janela_exibir = tk.Toplevel()
    janela_exibir.title("Exibição da Tabela")

    frame = tk.Frame(janela_exibir)
    frame.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)

    scroll_x = tk.Scrollbar(frame, orient=tk.HORIZONTAL)
    scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
    scroll_y = tk.Scrollbar(frame)
    scroll_y.pack(side=tk.RIGHT, fill=tk.Y)

    txt_tabela = tk.Text(frame, wrap=tk.NONE, xscrollcommand=scroll_x.set, yscrollcommand=scroll_y.set, width=160, height=40)
    txt_tabela.pack(fill=tk.BOTH, expand=True)
    scroll_x.config(command=txt_tabela.xview)
    scroll_y.config(command=txt_tabela.yview)

    cor_laranja_claro = "#FFCC99"
    cor_padrao = "#FFFFFF"

    for i, linha in enumerate(str(tabela_sem_colunas_vazias).split("\n")):
        if i % 2 == 0:
            txt_tabela.insert(tk.END, linha + "\n", ("laranja",))
        else:
            txt_tabela.insert(tk.END, linha + "\n", ("branco",))

    txt_tabela.tag_configure("laranja", background=cor_laranja_claro)
    txt_tabela.tag_configure("branco", background=cor_padrao)
    txt_tabela.config(state=tk.DISABLED)

# Função para criar uma nova linha
def criar_linha():
    janela_criar = tk.Toplevel()
    janela_criar.title("Criar Linha")

    campos = []

    for i, cabecalho in enumerate(cabecalhos_unicos):
        label = tk.Label(janela_criar, text=cabecalho)
        label.grid(row=i, column=0)
        campo = tk.Entry(janela_criar)
        campo.grid(row=i, column=1)
        campos.append(campo)

    def salvar_linha():
        nova_linha = [campo.get() for campo in campos]
        aba_ativa.append(nova_linha)
        tabela.save("GESTAO_DE_EXAMES_PERIODICOS.xlsx")
        messagebox.showinfo("Sucesso", "Linha criada com sucesso!")
        janela_criar.destroy()

    botao_salvar = tk.Button(janela_criar, text="Salvar", command=salvar_linha)
    botao_salvar.grid(row=len(cabecalhos_unicos), column=1)

# Função para alterar uma linha existente (selecionando pela primeira coluna)
def alterar_linha():
    janela_alterar = tk.Toplevel()
    janela_alterar.title("Alterar Linha")

    label_linha = tk.Label(janela_alterar, text="Valor da primeira coluna para alterar:")
    label_linha.grid(row=0, column=0)
    entrada_valor = tk.Entry(janela_alterar)
    entrada_valor.grid(row=0, column=1)

    campos = []
    for i, cabecalho in enumerate(cabecalhos_unicos):
        label = tk.Label(janela_alterar, text=cabecalho)
        label.grid(row=i+1, column=0)
        campo = tk.Entry(janela_alterar)
        campo.grid(row=i+1, column=1)
        campos.append(campo)

    def salvar_alteracao():
        valor_chave = entrada_valor.get()
        for linha in aba_ativa.iter_rows(min_row=5, max_col=9):
            if str(linha[0].value) == valor_chave:
                for i, campo in enumerate(campos):
                    aba_ativa.cell(row=linha[0].row, column=i+1, value=campo.get())
                tabela.save("GESTAO_DE_EXAMES_PERIODICOS.xlsx")
                messagebox.showinfo("Sucesso", "Linha alterada com sucesso!")
                janela_alterar.destroy()
                return
        messagebox.showerror("Erro", "Valor não encontrado!")

    botao_salvar = tk.Button(janela_alterar, text="Salvar Alterações", command=salvar_alteracao)
    botao_salvar.grid(row=len(cabecalhos_unicos)+1, column=1)

# Função para excluir uma linha existente (selecionando pela primeira coluna)
def excluir_linha():
    janela_excluir = tk.Toplevel()
    janela_excluir.title("Excluir Linha")

    label_linha = tk.Label(janela_excluir, text="Valor da primeira coluna para excluir:")
    label_linha.grid(row=0, column=0)
    entrada_valor = tk.Entry(janela_excluir)
    entrada_valor.grid(row=0, column=1)

    def confirmar_exclusao():
        valor_chave = entrada_valor.get()
        for linha in aba_ativa.iter_rows(min_row=5, max_col=9):
            if str(linha[0].value) == valor_chave:
                aba_ativa.delete_rows(linha[0].row)
                tabela.save("GESTAO_DE_EXAMES_PERIODICOS.xlsx")
                messagebox.showinfo("Sucesso", "Linha excluída com sucesso!")
                janela_excluir.destroy()
                return
        messagebox.showerror("Erro", "Valor não encontrado!")

    botao_excluir = tk.Button(janela_excluir, text="Excluir", command=confirmar_exclusao)
    botao_excluir.grid(row=1, column=1)

# Função para criar a tela inicial com os botões de ação
def tela_inicial():
    janela = tk.Tk()
    janela.title("Gestão de Exames Periódicos")

    botao_exibir = tk.Button(janela, text="Exibir Linhas", command=exibir_tabela)
    botao_exibir.pack(padx=10, pady=10)

    botao_criar = tk.Button(janela, text="Criar Linha", command=criar_linha)
    botao_criar.pack(padx=10, pady=10)

    botao_alterar = tk.Button(janela, text="Alterar Linha", command=alterar_linha)
    botao_alterar.pack(padx=10, pady=10)

    botao_excluir = tk.Button(janela, text="Excluir Linha", command=excluir_linha)
    botao_excluir.pack(padx=10, pady=10)

    janela.mainloop()

# Iniciar a tela inicial
tela_inicial()